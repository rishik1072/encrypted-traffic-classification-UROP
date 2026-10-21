"""
Test Suite: Supervisor Review Package and Research Paper Finalization.

Validates:
1. Existence, completeness, and structural integrity of docs/research/SUPERVISOR_EXECUTIVE_SUMMARY.md.
2. Existence and prioritization taxonomy of docs/research/NEXT_EXPERIMENTS_ROADMAP.md.
3. Existence and slide structure of docs/presentation/SUPERVISOR_PRESENTATION_GUIDE.md.
4. Consistent terminology, exact denominators, and limitation disclosures across all new package documents.
"""

from __future__ import annotations

from pathlib import Path
import unittest

PROJECT_ROOT = Path(__file__).resolve().parent.parent
DOCS_RESEARCH_DIR = PROJECT_ROOT / "docs" / "research"
DOCS_PRESENTATION_DIR = PROJECT_ROOT / "docs" / "presentation"

EXEC_SUMMARY_PATH = DOCS_RESEARCH_DIR / "SUPERVISOR_EXECUTIVE_SUMMARY.md"
ROADMAP_PATH = DOCS_RESEARCH_DIR / "NEXT_EXPERIMENTS_ROADMAP.md"
PRESENTATION_GUIDE_PATH = DOCS_PRESENTATION_DIR / "SUPERVISOR_PRESENTATION_GUIDE.md"


class TestSupervisorReviewPackage(unittest.TestCase):
    """Regression tests for the supervisor review package and handoff artifacts."""

    def test_supervisor_executive_summary_contains_all_required_sections(self) -> None:
        """Verify SUPERVISOR_EXECUTIVE_SUMMARY.md exists and includes sections 1 through 10."""
        self.assertTrue(EXEC_SUMMARY_PATH.exists(), f"Missing {EXEC_SUMMARY_PATH}")
        content = EXEC_SUMMARY_PATH.read_text(encoding="utf-8")

        required_sections = [
            "1. Research Problem and Motivation",
            "2. Project Objectives",
            "3. Dataset Composition and Provenance",
            "4. Experimental Methodology",
            "5. Main Evidence-Backed Findings",
            "6. Model Efficiency & Edge Deployment Footprint",
            "7. Important Scientific Limitations",
            "8. Unresolved Questions",
            "9. Recommended Next Experiments",
            "10. Specific Questions Requiring Supervisor Guidance",
        ]

        for sec in required_sections:
            self.assertIn(sec, content, f"Missing required section in executive summary: {sec}")

        # Core metric and disclosure verifications
        self.assertIn("READY WITH EXPLICIT LIMITATIONS", content)
        self.assertIn("0.5854", content)
        self.assertIn("0.5873", content)
        self.assertIn("0.5802", content)
        self.assertIn("0.0833", content)
        self.assertIn("28/48", content)
        self.assertIn("1/12", content)
        self.assertIn("UNRESOLVED", content)
        self.assertIn("HYBRID_PERTURBED", content)

    def test_experimental_roadmap_prioritization_and_category_separation(self) -> None:
        """Verify NEXT_EXPERIMENTS_ROADMAP.md cleanly separates Category A, Category B, and Category C."""
        self.assertTrue(ROADMAP_PATH.exists(), f"Missing {ROADMAP_PATH}")
        content = ROADMAP_PATH.read_text(encoding="utf-8")

        required_categories = [
            "Category A: Experiments Using Preserved Artifacts",
            "Category B: Retraining Experiments Using Existing Datasets",
            "Category C: Experiments Requiring New Physical Captures",
        ]

        for cat in required_categories:
            self.assertIn(cat, content, f"Roadmap missing category: {cat}")

        # Check key planned experiments
        self.assertIn("EXP-A1: Post-Hoc Conformal Calibration", content)
        self.assertIn("EXP-A2: C-Header / Bare-Metal Microcontroller Export", content)
        self.assertIn("EXP-C1: Multi-Device, Multi-Session Physical Wi-Fi Capture Campaign", content)
        self.assertIn("EXP-C2: Native Physical Ethernet & 4G/5G Cellular Hardware PCAPs", content)

    def test_presentation_guide_contains_ten_slides_with_required_elements(self) -> None:
        """Verify SUPERVISOR_PRESENTATION_GUIDE.md provides all 10 slides with messages, caveats, and notes."""
        self.assertTrue(PRESENTATION_GUIDE_PATH.exists(), f"Missing {PRESENTATION_GUIDE_PATH}")
        content = PRESENTATION_GUIDE_PATH.read_text(encoding="utf-8")

        for slide_num in range(1, 11):
            slide_header = f"## Slide {slide_num}:"
            self.assertIn(slide_header, content, f"Missing {slide_header} in presentation guide")

        # Verify key elements present throughout the presentation
        self.assertIn("Suggested Speaker Notes:", content)
        self.assertIn("Important Caveat:", content)
        self.assertIn("Main Message:", content)
        self.assertIn("Evidence / Visual to Use:", content)

    def test_all_referenced_figures_in_presentation_exist_on_disk(self) -> None:
        """Verify that every figure file referenced in the presentation guide actually exists in results/figures/."""
        figures_dir = PROJECT_ROOT / "results" / "figures"
        referenced_figures = [
            "research_baseline_model_comparison.png",
            "latency_vs_f1.png",
            "early_prediction_latency_vs_packets.png",
            "early_prediction_coverage_vs_packets.png",
            "coverage_vs_accuracy.png",
            "coverage_vs_error.png",
            "tunnel_shift.png",
            "environment_shift.png",
            "model_size_vs_f1.png",
            "resource_usage.png",
        ]

        for fig_name in referenced_figures:
            fig_path = figures_dir / fig_name
            self.assertTrue(fig_path.exists(), f"Referenced presentation figure missing on disk: {fig_path}")

    def test_urop_presentation_pptx_validity_and_embedded_notes(self) -> None:
        """Verify that UROP_Research_Presentation.pptx opens and contains 10 slides with non-empty speaker notes."""
        from pptx import Presentation

        pptx_path = DOCS_PRESENTATION_DIR / "UROP_Research_Presentation.pptx"
        self.assertTrue(pptx_path.exists(), f"Missing presentation file: {pptx_path}")

        prs = Presentation(str(pptx_path))
        self.assertEqual(len(prs.slides), 10, f"Expected 10 slides, got {len(prs.slides)}")

        for idx, slide in enumerate(prs.slides, 1):
            self.assertTrue(slide.has_notes_slide, f"Slide {idx} missing notes slide")
            notes = slide.notes_slide.notes_text_frame.text.strip()
            self.assertGreater(len(notes), 50, f"Slide {idx} notes are too short ({len(notes)} chars)")

    def test_urop_speaker_notes_markdown_coverage_and_defense_qa(self) -> None:
        """Verify that docs/presentation/UROP_Speaker_Notes.md exists and covers all 10 slides and defense Q&A."""
        speaker_notes_path = DOCS_PRESENTATION_DIR / "UROP_Speaker_Notes.md"
        self.assertTrue(speaker_notes_path.exists(), f"Missing speaker notes file: {speaker_notes_path}")
        content = speaker_notes_path.read_text(encoding="utf-8")

        for slide_num in range(1, 11):
            slide_header = f"## Slide {slide_num}:"
            self.assertIn(slide_header, content, f"Missing {slide_header} in speaker notes document")

        # Verify key elements
        self.assertIn("Speaking Script", content)
        self.assertIn("Slide Transition", content)
        self.assertIn("Anticipated Supervisor Question", content)
        self.assertIn("Evidence-Based Answer:", content)
        self.assertIn("Supervisor Defense Cheat-Sheet", content)


if __name__ == "__main__":
    unittest.main()
